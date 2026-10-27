"""Fill the retained Report 5 template; preserve every non-content package part."""
from pathlib import Path
from copy import deepcopy
import json, zipfile, hashlib, io
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path(__file__).resolve().parents[1]
ref=BASE/'work/templates/Report5_Test Documentation.docx'
assert hashlib.sha256(ref.read_bytes()).hexdigest()=='4ff653d4794fd23b905e9c85f26e33adf2ea6851cb15a1be3fa7f9270e4a2dfc'
data=json.loads((BASE/'work/report-data.json').read_text())
d=Document(ref);p=list(d.paragraphs);tables=list(d.tables)
pending='Đợi HuyLong điền'

def put(par,text,size=11):
    par.clear();r=par.add_run(text);r.font.name='Arial';r.font.size=Pt(size);r.font.color.rgb=RGBColor(0,0,0);r.italic=False
    return par
def after(par,text,style='Normal'):
    el=OxmlElement('w:p');par._p.addnext(el)
    from docx.text.paragraph import Paragraph
    new=Paragraph(el,par._parent);new.style=style
    if style=='Normal':put(new,text)
    else:new.add_run(text)
    return new
def bodyrows(t,rows,header=1):
    template=deepcopy(t.rows[-1]._tr)
    while len(t.rows)>header:t._tbl.remove(t.rows[-1]._tr)
    for row in rows:
        t._tbl.append(deepcopy(template));r=t.rows[-1]
        for cell,v in zip(r.cells,row):
            cell.text=str(v)
            for par in cell.paragraphs:
                for run in par.runs:run.font.name='Arial';run.font.size=Pt(10.5);run.font.color.rgb=RGBColor(0,0,0);run.italic=False
        trpr=r._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
    for i in range(header):t.rows[i]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))

# Existing cover spacing and images remain, metadata occupies previously blank slots.
put(p[11],'SmartLibrary',14)
put(p[12],pending+' — official project name and code',10)
put(p[13],pending+' — Report 5 author, issue date and version',10)
put(p[14],'Test evidence baseline: dev e5beb8e | Execution: 27 October 2026',10)
put(p[20],'Ho Chi Minh City — '+pending+' (official issue date)',10)
bodyrows(tables[0],[['2026-10-27','A',pending+' — document author','Added source-aligned test scope, preserved test results, local API measurements and open defect R5-DEF-001.']])

texts={
27:'This report documents the test scope, strategy, execution evidence and remaining risks for SmartLibrary. The baseline unit suite passed 203 cases; API integration passed 17, failed 1 and skipped 3. The failed registration-state transition remains open. These results do not establish release readiness or full mandatory-feature acceptance.',
28:'Scope covers identity and external-reader review; catalog and copies; borrowing, returns, incidents, reservations and renewals; fines and payment; profiles, reviews and support; news and notifications; user administration, policies, AI configuration/search and reporting. The companion workbook groups these into 16 candidate mandatory features and maps all 40 Report 3 use cases. Grouping does not replace the signed Review 2 commitment.',
29:'Code is the reference for implemented behavior; approved requirements remain the reference for commitments. Report 3 and Report 4 are still incomplete. Conflicts must be resolved explicitly. Authentication and external providers are doubled in API tests. Production migrations, real Supabase login, real PayOS/signatures, email delivery, Cloudinary, browser workflows, AI-provider latency and sustained production capacity remain unverified.',
31:'Use risk-based testing, prioritizing circulation integrity, registration transitions, fine settlement and authorization. Preserve failed and skipped results. Run targeted regressions after changes rather than repeatedly rerunning every suite. Separate unit, integration, local measurement, browser system and user-acceptance evidence.',
33:'Functional tests cover happy paths, input validation and permission failures. State-transition tests cover valid and invalid entity transitions. Concurrency tests target the last-copy race and duplicate settlement. Performance tests record actual local HTTP latency/error samples. Completion requires the applicable expected results and persisted effects to agree; no passing result may be inferred from a mock for a real service.',
35:'Unit tests isolate rules and services using fixtures/doubles. API integration executes real controllers, services and PostgreSQL transactions with test authentication. System/E2E tests must verify complete actor workflows in a browser. Acceptance testing requires HuyLong and the designated reviewer to confirm the committed scope; it has not been performed.',
38:'Versioned test code, raw TRX and coverage accompany this document. The load experiment is an opt-in xUnit probe, not a general performance certification.',
42:'HuyLong is coordinating the remaining testing and documentation. Automation supports execution and evidence collection; it does not replace human review or final acceptance. Actual team assignments below must be confirmed.',
45:'The isolated fixture creates a random localhost database and removes it on disposal; it does not load repository .env or modify the shared application database. It uses EnsureCreated and text domains, not the full Supabase/pgvector production migration chain. Test accounts and records are synthetic.',
48:'The sequence is baseline/evidence preservation, local measurement, traceability and document preparation, then remaining system/provider tests, defect retest and acceptance. Official review dates have not been supplied.',
51:'The accompanying workbooks contain executable-source references, expected/actual results and evidence locators. Temporary IDs are used only for this report package. '+pending+' — confirm the official requirement and test ID convention before submission.',
52:'Unit cases: Report5_Unit Test.xlsx — 203 recorded executions organized into 93 related function groups. Conditions, theory inputs, assertion excerpts and source locations are included. N/A/B categories are provisional; full assertions remain in the referenced source.',
53:'Other cases: Report5_Test Report.xlsx — 87 entries: 21 integration cases (including three skipped placeholders), four local NFR scenarios and 62 planned system/E2E/state scenarios. It also contains the feature inventory, workflows, 40-use-case RTM, unit traceability and defect register.',
55:'The evidence supports a working regression baseline, not a release sign-off. Unit results are 203/203 passed. Executed API integration results are 17/18 passed (94.4%), with three further cases skipped. Planned system/browser cases are not included in that denominator. One local load probe collected four concurrency scenarios; its successful completion is separate from business-test results.'}
for idx,text in texts.items():put(p[idx],text)

bodyrows(tables[1],[['Functional / validation','Executed','Executed','Planned','Planned'],['State / authorization','Executed in part','Executed in part','Planned','Planned'],['Concurrency / performance','Rules only','Partial / measured locally','Not executed','Not executed']],2)
bodyrows(tables[2],[['Unit / integration','xUnit, Moq, TestHost','Open source','2.9.3 / 4.21.0 / 10.0.0'],['Coverage / evidence','coverlet collector; TRX','Open source','6.0.4'],['Build and CI','.NET SDK; GitHub Actions','Microsoft / GitHub','10.0.402; ci.yml'],['API load probe','Kestrel + HttpClient + Stopwatch','Project test harness','.NET runtime 10.0.13']])
bodyrows(tables[3],[['HuyLong','Coordinator / reviewer','Confirm scope, metadata and final evidence; approve outstanding testing.'],[pending,'Actual team roles',pending+' — list members and their real responsibilities; do not infer from commit authors.'],['Automation','Test execution support','Run isolated checks, preserve evidence and prepare reports; no independent acceptance authority.']])
bodyrows(tables[4],[['Local OS / hardware','macOS Arm64, 10 logical processors','Local machine','15.3; CPU model/RAM unavailable'],['API / database','Kestrel / PostgreSQL on 127.0.0.1:55437','Local isolated environment','10.0.13 / 17.11'],['Load dataset','1000 titles, 4000 copies, 3 accounts, 0 loans','Synthetic seed','NfrProbeTests'],['Identity / providers','Test role headers; PayOS, email and storage doubles','Test-only adapters','No live-provider verification']])
bodyrows(tables[5],[['Baseline and isolated tests','2026-10-27','2026-10-27'],['Local API measurements and report preparation','2026-10-27','2026-10-27'],['Provider/browser tests and defect retest',pending+' — planned start',pending+' — planned end'],['Review 2 / Defense 1.1 / Report 5 submission',pending+' — official milestones',pending+' — submission deadline']])

anchor=p[55]
sections=[
('5.1 Evidence and acceptance','Heading 3'),
('Evidence paths are relative to docs/report5: evidence/unit.trx, evidence/integration.trx, evidence/unit-coverage.cobertura.xml and evidence/nfr-local.json. The successful measurement TRX is under evidence/nfr-verified. The initial NFR warmup failure was a test-fixture configuration issue and is not a product defect. The fixture was corrected without changing business code.','Normal'),
('The saved functional results precede the network-mode fixture extension; the full suite was not rerun. The new CI configuration is local and has not been pushed or executed remotely. Existing remote green runs are not proof that the new regression passes. Frontend lint/build previously exited successfully, but their console output was not archived as a separate artifact.','Normal'),
('The Review 2 guide sets Defense 1.1 targets, including at least three cases per mandatory feature, one E2E per main workflow, state-transition and NFR coverage, at least 90% overall passing, 100% mandatory main-flow passing, at least 50% automated in CI and no open Critical/High defects. These are targets, not claims of achievement. This package is not eligible for acceptance while scope, unexecuted workflows and the proposed-High defect remain unresolved.','Normal'),
('5.2 Local performance measurements','Heading 3'),
('The experiment ran HTTP loopback Kestrel and PostgreSQL on the same macOS 15.3 Arm64 machine with 10 logical processors. CPU model and RAM could not be read because hardware-query permission was denied. No TLS, browser rendering, real authentication or provider calls were included. The three read endpoints were catalog page 1, catalog page 10 (20 books per page), and copy summary, in an equal mix.','Normal'),
('After 60 warmup requests, 1200 requests were issued at each maximum concurrency level 1, 10, 50 and 100 using a closed loop with no think time. The load generator shared the server process. Percentiles use the nearest-rank method and include response body reads; error counts include non-success responses and request exceptions. Only one short run per level was measured; warm-cache and same-host effects limit generalization.','Normal'),
('At concurrency 1/10/50/100, P95 was 3.73/13.22/50.08/152.29 ms and P99 was 4.76/22.41/85.65/195.97 ms. All 4800 responses succeeded. Measured throughput was 320.4/1351.1/1360.1/1135.9 requests per second. Each burst lasted only 0.88–3.75 seconds. These figures are not evidence of sustained 100-user capacity.','Normal'),
('Đề xuất: use API P95 below 2000 ms and zero HTTP errors as a local screening criterion. The four bursts meet that proposed criterion. Report 3 separately states standard page loads under 2 seconds, AI under 3.5 seconds and 100 active users; the local API experiment does not certify those end-to-end targets. Availability 99.9%, MTBF 180 days, MTTR 2 hours, 50000 titles/200000 copies and CPU/RAM below 70% remain Not Tested.','Normal'),
('5.3 Open defect','Heading 3'),
('R5-DEF-001 — Approved registration accepts request-more-info. Status: Open / Fail. Severity: High proposed, pending confirmation. Seed a Submitted request, approve it as Librarian, then POST request-more-info for the same ID with a reason. Expected: reject the invalid transition and retain Approved/Active. Actual: the second call succeeds, failing Registration_ApprovedRequest_CannotBeReopenedForDocuments.','Normal'),
('RegistrationService.RequestMoreInfoAsync (baseline lines 193–208) checks existence but not current status, then assigns NeedMoreInfo and PendingDocuments and saves. Those assignments are confirmed by source inspection. The failed test stops at the HTTP assertion, before its subsequent database assertion; do not claim a separately observed post-action DB check. No business fix has been applied.','Normal'),
('5.4 Traceability and remaining risks','Heading 3'),
('The RTM maps UC01–UC40 to candidate feature groups, implementation references, test IDs and saved evidence. UC18 is Fail due to R5-DEF-001. Other use cases remain Not Tested at full-requirement acceptance level because only parts of their behavior have executed evidence. Per-case Passed results remain visible. Unit traceability is supporting module-level evidence, not a claim that an entire UC is verified.','Normal'),
('Important reconciliation items: Report 3 names payment providers other than the implemented PayOS; its review eligibility rule requires prior borrowing, while the implementation needs confirmation against that rule; UC29 import-ticket CRUD must not be declared implemented merely because copy intake exists. Realtime endpoint, authentication-storage and NFR claims also require alignment. Do not silently change a requirement to make the current code pass.','Normal'),
('The nine planned browser workflows cover registration, borrowing/handover, returns/incidents, fines, catalog/copies, reservation/renewal, support/news, administration/reporting and AI discovery. Their detailed steps are in the Main Workflows sheet. Browser E2E and live-provider tests remain Not Tested in this delivery. No real payment or notification delivery was triggered by the isolated test runs.','Normal'),
('Before submission, '+pending+' — fill official project name/code, author, issue date, version, members/roles, Review 2 and Defense 1.1 dates, submission deadline, requirement-ID convention and the official evidence archive link. Confirm the signed mandatory scope and proposed defect severity. Preserve Fail/Not Tested statuses until new evidence justifies a change.','Normal')]
for text,style in sections:anchor=after(anchor,text,style)
# Preserve the template's separate TOC page even when cover slots change height.
for el in d._element.iter(qn('w:p')):
    if ''.join(t.text or '' for t in el.iter(qn('w:t')))=='Table of Contents':
        pr=el.find(qn('w:pPr'))
        if pr is None:pr=OxmlElement('w:pPr');el.insert(0,pr)
        pr.append(OxmlElement('w:pageBreakBefore'))
# Rebalance only filled tables whose placeholder column widths forced words apart.
for ti,widths in [(3,[1.35,1.35,3.6]),(4,[1.3,2.25,1.25,1.5]),(5,[3.4,1.45,1.45])]:
    t=tables[ti];t.autofit=False
    for c,width in zip(t.columns,widths):c.width=Inches(width)
    for row in t.rows:
        for cell,width in zip(row.cells,widths):cell.width=Inches(width)
# Retain fields and request refresh when opened in Word.
upd=d.settings.element.find(qn('w:updateFields'))
if upd is None:upd=OxmlElement('w:updateFields');d.settings.element.append(upd)
upd.set(qn('w:val'),'true')
for t in d.tables:
    for r in t.rows:
        for cell in r.cells:
            for par in cell.paragraphs:par.paragraph_format.keep_with_next=False
buf=io.BytesIO();d.save(buf);buf.seek(0)
out=BASE/'deliverables/Report5_Test Documentation.docx';out.parent.mkdir(exist_ok=True)
with zipfile.ZipFile(ref) as z, zipfile.ZipFile(buf) as modified,zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as final:
    for info in z.infolist():final.writestr(info,modified.read(info.filename) if info.filename in ('word/document.xml','word/settings.xml') else z.read(info.filename))
with zipfile.ZipFile(ref) as a,zipfile.ZipFile(out) as b:
    assert all(a.read(n)==b.read(n) for n in a.namelist() if n not in ('word/document.xml','word/settings.xml'))
print(out)
