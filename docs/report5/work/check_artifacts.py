"""Read-only numeric/text audit. No tests, database access, rendering or file writes."""
from pathlib import Path
from collections import Counter
import json
import re
import zipfile
import xml.etree.ElementTree as ET
import openpyxl

BASE = Path(__file__).resolve().parents[1]
data = json.loads((BASE / 'work/report-data.json').read_text())
evidence = json.loads((BASE / 'evidence/nfr-local.json').read_text())
issues = []
def check(ok, message):
    print(('OK: ' if ok else 'FAIL: ') + message)
    if not ok: issues.append(message)

counts = Counter(c['status'] for c in data['cases'])
check(len(data['cases']) == 87 and counts == {'Passed':21, 'Failed':1, 'Not Tested':65}, f'Non-unit inventory: {len(data["cases"])} {dict(counts)}')
check(len({c['id'] for c in data['cases']}) == len(data['cases']), 'Unique non-unit case IDs')
check(len(data['unit']) == 203 and len({c['group'] for c in data['unit']}) == 93, '203 unit results / 93 function groups')
check(data['nfr'] == evidence, 'NFR data equals saved evidence')

books = {}
texts = {}
placeholders = {}
for name in ['Report5_Test Report.xlsx', 'Report5_Unit Test.xlsx']:
    path = BASE / 'deliverables' / name
    values = openpyxl.load_workbook(path, data_only=True)
    formulas = openpyxl.load_workbook(path, data_only=False)
    books[name] = (values, formulas)
    text = []
    errors = []
    unexpected = []
    missing_cache = []
    slots = []
    for sheet in values:
        for row in sheet:
            for c in row:
                f = formulas[sheet.title][c.coordinate]
                v = c.value
                if v is not None: text.append(str(v))
                if c.data_type == 'e' or (isinstance(v,str) and re.search(r'#REF!|#VALUE!|#NAME\?',v)):
                    errors.append(f'{sheet.title}!{c.coordinate}')
                if f.data_type == 'f':
                    if v is None: missing_cache.append(f'{sheet.title}!{c.coordinate}')
                    if not re.match(r'=(?:COUNTIF\(|COUNTA\(|SUM\(|IFERROR\(|Cover!|\x27)',f.value):
                        unexpected.append(f'{sheet.title}!{c.coordinate}')
                if isinstance(v,str) and 'Đợi HuyLong điền' in v:
                    slots.append((f'{sheet.title}!{c.coordinate}',v))
    texts[name] = '\n'.join(text)
    placeholders[name] = slots
    check(not errors, f'{name}: no error cells/strings ({errors})')
    check(not unexpected, f'{name}: no unintended formulas ({unexpected})')
    check(not missing_cache, f'{name}: formula result caches present ({missing_cache})')

report = books['Report5_Test Report.xlsx'][0]
stat = report['Test Statistics']
check([stat[a].value for a in ['D14','E14','F14','G14','H14']] == [21,1,65,0,87], 'Test Report statistics reconcile to JSON')
observed = {}
for name in ['Feature 1','Feature 2']:
    for row in report[name].iter_rows(min_row=12,values_only=True):
        if row[0]: observed[row[0]] = row[5]
check(observed == {c['id']:c['status'] for c in data['cases']}, 'All non-unit case verdicts match JSON')
for i,s in enumerate(evidence['summaries'],5):
    expected = [s[k] for k in ['concurrency','requests','errors','seconds','requestsPerSecond','meanMs','p50Ms','p95Ms','p99Ms','maxMs']]
    actual = [report['NFR Results'].cell(i,c).value for c in range(1,11)]
    check(all(abs(a-b)<1e-8 for a,b in zip(actual,expected)), f'NFR concurrency {s["concurrency"]}: all metrics match, P95={s["p95Ms"]}')

uv, uf = books['Report5_Unit Test.xlsx']
function_sheets = [s for s in uv.sheetnames if re.fullmatch(r'F\d{3}',s)]
check(len(function_sheets)==93, 'Unit workbook has 93 function matrices')
check([uv['Statistics'].cell(99,c).value for c in range(2,6)] == [203,0,0,203], 'Unit statistics: 203 Passed / 0 Failed / 0 Not Tested')
unit_by_id = {r['id']:r for r in data['unit']}
seen = []
literal_count = 0
for name in function_sheets:
    s = uv[name]
    for c in range(6,s.max_column+1):
        uid = s.cell(9,c).value
        if not uid: continue
        seen.append(uid)
        expected = unit_by_id[uid]['expected']
        cell = uf[name].cell(13,c)
        actual = cell.value
        # The exporting library may retain the text-escape apostrophe.
        if isinstance(actual,str) and actual.startswith("'="): actual=actual[1:]
        check(cell.data_type != 'f' and actual == expected, f'{name}!{cell.coordinate}: assertion text preserved') if (cell.data_type == 'f' or actual != expected) else None
        if expected.startswith('=>'): literal_count += 1
        if s.cell(18,c).value != unit_by_id[uid]['status']: issues.append(f'{name}: verdict mismatch')
check(set(seen)==set(unit_by_id) and len(seen)==203, 'All 203 unit IDs and verdicts reconcile')
check(literal_count>0, f'Leading => assertion excerpts checked as literal text: {literal_count}')

docpath = BASE / 'deliverables/Report5_Test Documentation.docx'
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with zipfile.ZipFile(docpath) as z:
    root=ET.fromstring(z.read('word/document.xml'))
    paragraphs=[''.join(t.text or '' for t in p.findall('.//w:t',ns)) for p in root.findall('.//w:p',ns)]
text='\n'.join(paragraphs)
texts[docpath.name]=text
placeholders[docpath.name]=[(f'paragraph {i+1}',p) for i,p in enumerate(paragraphs) if 'Đợi HuyLong điền' in p]
check('87 entries' in text and re.search(r'21 Passed\s*/\s*1 Failed\s*/\s*65 Not Tested',text), 'DOCX: non-unit total and verdict breakdown reconcile')
for s in evidence['summaries']:
    check(f'{s["p95Ms"]:.2f}' in text,f'DOCX: P95 {s["p95Ms"]:.2f} present')
for name,text in texts.items():
    if name=='Report5_Test Report.xlsx':
        defect_ok=report['Defects']['A5'].value=='R5-DEF-001' and report['Defects']['B5'].value=='Open / Fail'
    else: defect_ok=bool(re.search(r'R5-DEF-001[^\n]*Fail',text))
    check(defect_ok,f'{name}: R5-DEF-001 explicitly Fail')
    check(not re.search(r'#REF!|#VALUE!|#NAME\?',text),f'{name}: no requested error strings')
    print('CLAIM REVIEW:',name)
    for line in dict.fromkeys(text.splitlines()):
        if re.search(r'acceptance|99\.9|100.user|bursts meet|certif|fully verified|feature.*pass|pass.*feature',line,re.I):
            print(' ',line)
print('PLACEHOLDERS (deduplicated by text):')
for name,slots in placeholders.items():
    grouped={}
    for loc,text in slots: grouped.setdefault(text,[]).append(loc)
    print(name)
    for text,locations in grouped.items():print(' ',locations[0],f'({len(locations)} occurrence(s)):',text)
print('AUDIT ISSUES:',issues)
raise SystemExit(bool(issues))
